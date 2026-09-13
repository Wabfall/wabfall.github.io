"""Dossier de compétences au format ESN — Etienne C., Data Engineer.

Usage : python generate_dossier.py   (dépendance : pip install python-docx)
Écrit Dossier_competences_Etienne_C_Data_Engineer.docx à côté de ce script.

Les faits viennent de src/data/portfolio.ts, du rapport TOEIC (ETS) et du badge Credly.
N'y ajouter une compétence que si elle a été pratiquée : un dossier ESN se vérifie en
entretien technique. Pas de TJM ni de coordonnées ici : l'ESN les gère elle-même.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(__file__).with_name("Dossier_competences_Etienne_C_Data_Engineer.docx")

NAVY = RGBColor(0x0B, 0x2E, 0x4F)
ACCENT = RGBColor(0x1F, 0x6F, 0xB2)
GREY = RGBColor(0x55, 0x5E, 0x68)
INK = RGBColor(0x1E, 0x25, 0x2D)
FONT = "Calibri"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.top_margin, sec.bottom_margin = Cm(1.5), Cm(1.5)

normal = doc.styles["Normal"]
normal.font.name = FONT
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.font.size = Pt(10)
normal.font.color.rgb = INK
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.line_spacing = 1.1


def shade(cell, hex_fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tcPr.append(shd)


def cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    mar = OxmlElement("w:tblCellMar")
    for side, val in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        el = OxmlElement(f"w:{side}")
        el.set(qn("w:w"), str(val))
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)


def borders(table, color="D5DCE3", inside=True):
    tblPr = table._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    edges = ["top", "left", "bottom", "right"] + (["insideH", "insideV"] if inside else [])
    for edge in ["top", "left", "bottom", "right", "insideH", "insideV"]:
        el = OxmlElement(f"w:{edge}")
        if edge in edges:
            el.set(qn("w:val"), "single")
            el.set(qn("w:sz"), "4")
            el.set(qn("w:color"), color)
        else:
            el.set(qn("w:val"), "nil")
        b.append(el)
    tblPr.append(b)


def set_widths(table, widths):
    """Fixe la grille : sans elle, Word et LibreOffice répartissent les colonnes à parts égales."""
    table.autofit = False
    grid = table._tbl.tblGrid
    for col, w in zip(grid.findall(qn("w:gridCol")), widths):
        col.set(qn("w:w"), str(int(w.twips)))
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w


def keep_with_next(p):
    p.paragraph_format.keep_with_next = True


def run(p, text, bold=False, size=None, color=None, italic=False):
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return r


def para(text="", size=10, color=INK, bold=False, before=0, after=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align:
        p.alignment = align
    if text:
        run(p, text, bold=bold, size=size, color=color)
    return p


def section_title(text):
    p = para(before=12, after=5)
    run(p, text.upper(), bold=True, size=12, color=NAVY)
    pPr = p._p.get_or_add_pPr()
    bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "2")
    bottom.set(qn("w:color"), "1F6FB2")
    bdr.append(bottom)
    pPr.append(bdr)
    keep_with_next(p)
    return p


def bullet(text, lead=None, size=10):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.left_indent = Cm(0.6)
    p.paragraph_format.first_line_indent = Cm(-0.4)
    if lead:
        run(p, lead, bold=True, size=size)
    run(p, text, size=size)
    return p


def label(text):
    p = para(before=5, after=2)
    run(p, text, bold=True, size=10, color=ACCENT)
    keep_with_next(p)
    return p


# ─── En-tête ────────────────────────────────────────────────────────────────
head = doc.add_table(rows=1, cols=1)
head.alignment = WD_TABLE_ALIGNMENT.CENTER
borders(head, inside=False, color="0B2E4F")
cell_margins(head, top=180, bottom=180, left=260, right=260)
c = head.cell(0, 0)
shade(c, "0B2E4F")
p = c.paragraphs[0]
run(p, "ETIENNE C.", bold=True, size=22, color=RGBColor(0xFF, 0xFF, 0xFF))
p = c.add_paragraph()
run(p, "Data Engineer", bold=True, size=14, color=RGBColor(0xFF, 0xFF, 0xFF))
run(p, "   ·   GCP · BigQuery · dbt · Python", size=12, color=RGBColor(0xB9, 0xD4, 0xEE))
p = c.add_paragraph()
p.paragraph_format.space_before = Pt(4)
run(p, "Double compétence data engineering et génie logiciel", italic=True, size=10,
    color=RGBColor(0xDD, 0xE7, 0xF1))

para(after=4)

facts = [
    ("Expérience", "4 ans, dont 3 en alternance"),
    ("Disponibilité", "Janvier 2027"),
    ("Mobilité", "Grand Ouest en priorité, ouvert sur toute la France"),
    ("Mode de travail", "Hybride ou full remote, présentiel possible"),
    ("Langues", "Français natif · Anglais professionnel (TOEIC 880/990)"),
    ("Certification", "Google Cloud Digital Leader"),
]
ft = doc.add_table(rows=len(facts), cols=2)
borders(ft)
cell_margins(ft, top=40, bottom=40)
for i, (k, v) in enumerate(facts):
    kc, vc = ft.cell(i, 0), ft.cell(i, 1)
    kc.width, vc.width = Cm(3.6), Cm(13.8)
    shade(kc, "EEF3F8")
    run(kc.paragraphs[0], k, bold=True, size=9.5, color=NAVY)
    run(vc.paragraphs[0], v, size=9.5)
set_widths(ft, (Cm(3.6), Cm(13.8)))

# ─── Synthèse ───────────────────────────────────────────────────────────────
section_title("Synthèse du profil")
p = para(after=4)
run(p, "Ingénieur ESIEA, trois ans d'alternance chez Bouygues Telecom entre développement applicatif "
       "et équipe data, puis Data Engineer chez papernest à Barcelone : seul Data Engineer sur site pour "
       "les marchés français, espagnol et italien, en charge du moteur ETL interne en Python et des "
       "modèles dbt sur BigQuery, avec les standards du génie logiciel (tests, documentation, CI/CD). "
       "Recherche une mission de construction ou de fiabilisation de plateforme data, idéalement sur GCP.")

label("Points forts")
bullet("pipelines Python et dbt sur BigQuery, et outils internes full-stack (React, Next.js, Spring Boot) pour les équipes métier.",
       lead="Double compétence : ")
bullet("data contracts, validation de schéma, tests dbt et alerting pour détecter les anomalies en amont.",
       lead="Fiabilité de la donnée : ")
bullet("seul Data Engineer sur site, interlocuteur direct des équipes Sales Ops et Partnership sur trois marchés.",
       lead="Autonomie et relation métier : ")
bullet("conception d'outils internes pensés pour les usages réels des équipes, techniques ou non : éditeur "
       "de configuration des pipelines, outils à base de LLM.",
       lead="Outillage au service des métiers : ")

# ─── Compétences ────────────────────────────────────────────────────────────
section_title("Compétences techniques")
skills = [
    ("Data engineering", None, None),
    ("SQL, BigQuery", "Confirmé", "Bouygues Telecom puis papernest, depuis 2022"),
    ("dbt (modélisation, tests, documentation)", "Confirmé", "papernest, au quotidien"),
    ("Python (moteur ETL, outillage)", "Confirmé", "papernest, au quotidien"),
    ("Data contracts, validation de schéma, alerting", "Confirmé", "papernest"),
    ("Apache Flink SQL (streaming)", "Opérationnel", "Bouygues Telecom (3 ans) puis papernest (1 an)"),
    ("Teradata, Hadoop", "Opérationnel", "Bouygues Telecom, scripts de gouvernance"),
    ("ElasticSearch, MongoDB", "Opérationnel", "Bouygues Telecom, plateforme d'analyse de logs"),
    ("Cloud et DevOps", None, None),
    ("Google Cloud Platform", "Opérationnel", "BigQuery en production · certifié Cloud Digital Leader"),
    ("Git, CI/CD, GitHub Actions", "Confirmé", "Bouygues Telecom puis papernest, depuis 2022"),
    ("Docker, Linux", "Opérationnel", "papernest · services auto-hébergés sur VPS"),
    ("Développement et IA", None, None),
    ("TypeScript, React, Next.js", "Confirmé", "papernest (éditeur YAML), Bouygues Telecom"),
    ("Java, Spring Boot", "Opérationnel", "Bouygues Telecom (3 ans), stage Ag2ir"),
    ("APIs LLM, prompt engineering", "Opérationnel", "outils internes papernest"),
]
st = doc.add_table(rows=1, cols=3)
borders(st)
cell_margins(st, top=35, bottom=35)
widths = (Cm(6.6), Cm(2.6), Cm(8.2))
for i, h in enumerate(("Compétence", "Niveau", "Contexte de pratique")):
    hc = st.rows[0].cells[i]
    hc.width = widths[i]
    shade(hc, "0B2E4F")
    run(hc.paragraphs[0], h, bold=True, size=9.5, color=RGBColor(0xFF, 0xFF, 0xFF))
for name, level, ctx in skills:
    row = st.add_row()
    cells = row.cells
    if level is None:
        merged = cells[0].merge(cells[2])
        shade(merged, "EEF3F8")
        run(merged.paragraphs[0], name, bold=True, size=9.5, color=NAVY)
        continue
    for i, txt in enumerate((name, level, ctx)):
        cells[i].width = widths[i]
        run(cells[i].paragraphs[0], txt, size=9.5,
            bold=(i == 1), color=(ACCENT if i == 1 else INK))
set_widths(st, widths)
p = para(before=3)
run(p, "Confirmé : pratique autonome en production · Opérationnel : pratique réelle en mission ou en projet",
    italic=True, size=8.5, color=GREY)


# ─── Expériences ────────────────────────────────────────────────────────────
def experience(role, company, meta, period, context, achievements, env):
    t = doc.add_table(rows=1, cols=2)
    borders(t, inside=False, color="FFFFFF")
    cell_margins(t, top=60, bottom=60, left=120, right=120)
    l, r = t.cell(0, 0), t.cell(0, 1)
    l.width, r.width = Cm(12.4), Cm(5.0)
    shade(l, "EEF3F8")
    shade(r, "EEF3F8")
    p = l.paragraphs[0]
    run(p, role, bold=True, size=11.5, color=NAVY)
    run(p, "  |  " + company, bold=True, size=11.5, color=ACCENT)
    p = l.add_paragraph()
    run(p, meta, size=9, color=GREY)
    set_widths(t, (Cm(12.4), Cm(5.0)))
    p = r.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    run(p, period, bold=True, size=10, color=NAVY)
    t.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))

    label("Contexte")
    para(context, after=2)
    label("Réalisations")
    for a in achievements:
        if isinstance(a, tuple):
            bullet(a[1], lead=a[0])
        else:
            bullet(a)
    p = para(before=4, after=8)
    run(p, "Environnement technique : ", bold=True, size=9.5, color=ACCENT)
    run(p, env, size=9.5)


section_title("Expériences professionnelles")

experience(
    "Data Engineer", "papernest",
    "Barcelone, Espagne · Hybride · VIE",
    "10/2025 – aujourd'hui",
    "Scale-up qui simplifie la gestion des contrats du quotidien (énergie, box internet, assurance) "
    "sur les marchés français, espagnol et italien. Seul Data Engineer sur le site de Barcelone, dans "
    "une stack data en transition d'un fonctionnement historique vers une plateforme fiable.",
    [
        ("Moteur ETL et dbt : ", "développement de fonctionnalités sur le moteur ETL interne (Python) et "
         "sur les modèles dbt alimentant BigQuery, avec tests, documentation et CI/CD."),
        ("Qualité de la donnée : ", "mise en place de data contracts et de validation de schéma pour "
         "détecter les anomalies avant qu'elles n'atteignent les usages aval."),
        ("Supervision : ", "processus d'alerting sur les pipelines pour une détection proactive des incidents."),
        ("Éditeur de configuration YAML : ", "conception et développement d'un outil (Next.js, React) qui "
         "sécurise la configuration des flux, y compris pour les équipes non techniques. Le temps de "
         "fabrication et de test d'un pipeline a été divisé par deux environ, et l'investigation d'un "
         "incident est passée de 30 à 60 minutes à moins de 2 minutes."),
        ("Outils IA : ", "outils internes à base de LLM pour accélérer les workflows des équipes techniques "
         "et non techniques."),
        ("Rôle transverse : ", "interlocuteur data principal des équipes Sales Ops et Partnership sur les "
         "trois marchés."),
    ],
    "Python, dbt, BigQuery, SQL, Apache Flink, GCP, data contracts, Docker, YAML, Next.js, React, "
    "TypeScript, Tailwind CSS, Git, CI/CD, APIs LLM.",
)

experience(
    "Développeur Front-Data (alternance)", "Bouygues Telecom",
    "Nantes, France · Hybride",
    "09/2022 – 08/2025",
    "Alternance de trois ans en parallèle du cycle ingénieur, au sein de l'équipe data d'un opérateur "
    "télécom national : plateforme interne d'analyse de logs, ETL central et outillage de gouvernance.",
    [
        ("Plateforme d'analyse de logs : ", "développement full-stack d'une application web de "
         "visualisation et d'analyse des logs machines. Temps d'investigation ramené de plusieurs minutes "
         "à un accès quasi instantané."),
        ("Streaming : ", "amélioration d'un module SQL de streaming de l'ETL interne sous Apache Flink."),
        ("Gouvernance des données : ", "développement de scripts sur BigQuery, Teradata et Hadoop."),
        ("Documentation : ", "génération automatique de la documentation depuis le code Java et le "
         "Markdown avec Docusaurus."),
    ],
    "Java, Spring Boot, React, ElasticSearch, MongoDB, Apache Flink (SQL), BigQuery, Teradata, "
    "Hadoop, Docusaurus, Git.",
)

experience(
    "Développeur Java (stage)", "Ag2ir",
    "Alençon, France",
    "04/2022 – 06/2022",
    "Stage de fin de DUT : outillage autour de l'ERP EBP.",
    [
        "Développement d'une application de configuration de l'ERP à partir de fichiers XML, avec une "
        "interface Java Swing.",
        "Automatisation du traitement des fichiers XML pour faciliter l'intégration de l'ERP chez les clients.",
    ],
    "Java, Swing, XML.",
)

# ─── Formation, certifications, langues ─────────────────────────────────────
section_title("Formation et certifications")
edu = [
    ("2022 – 2025", "Diplôme d'ingénieur en informatique et technologies numériques, ESIEA (Laval), "
                    "en alternance. Majeure génie logiciel : full-stack, architecture cloud, CI/CD, Docker."),
    ("2023", "Semestre d'échange à l'Université du Québec à Chicoutimi (UQAC), Canada."),
    ("2020 – 2022", "DUT Informatique, IUT du Havre."),
    ("2026", "TOEIC Listening & Reading : 880/990 (niveau B2)."),
    ("2025", "Google Cloud Digital Leader (valide jusqu'en 2028)."),
]
et = doc.add_table(rows=len(edu), cols=2)
borders(et, inside=False, color="FFFFFF")
cell_margins(et, top=25, bottom=25, left=0, right=80)
for i, (d, t) in enumerate(edu):
    et.cell(i, 0).width, et.cell(i, 1).width = Cm(2.6), Cm(14.8)
    run(et.cell(i, 0).paragraphs[0], d, bold=True, size=10, color=ACCENT)
    run(et.cell(i, 1).paragraphs[0], t, size=10)
set_widths(et, (Cm(2.6), Cm(14.8)))

section_title("Langues")
bullet("langue maternelle.", lead="Français : ")
bullet("professionnel, langue de travail au quotidien chez papernest (TOEIC 880/990).", lead="Anglais : ")

# ─── Pied de page ───────────────────────────────────────────────────────────
fp = sec.footer.paragraphs[0]
fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
run(fp, "Etienne C. · Data Engineer · Dossier de compétences · page ", size=8, color=GREY)
r = fp.add_run()
r.font.size = Pt(8)
r.font.color.rgb = GREY
for tag, text in (("begin", None), (None, "PAGE"), ("end", None)):
    if tag:
        fc = OxmlElement("w:fldChar")
        fc.set(qn("w:fldCharType"), tag)
        r._r.append(fc)
    else:
        it = OxmlElement("w:instrText")
        it.set(qn("xml:space"), "preserve")
        it.text = text
        r._r.append(it)

doc.core_properties.title = "Dossier de compétences — Etienne C., Data Engineer"
doc.core_properties.author = "Etienne C."
doc.save(OUT)
print("écrit :", OUT)
